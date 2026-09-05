#!/usr/bin/env python3
"""Normalize the PartyReport workbook into an import-safe JSON document.

The workbook is treated strictly as input data. This script performs no network
or database operations.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


GSTIN_RE = re.compile(r"^\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]$")
EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
PINCODE_RE = re.compile(r"(?<!\d)([1-9]\d{5})(?!\d)")
PLACEHOLDERS = {"", "-", "0", "4", "N/A", "NA", "NONE", "NULL"}

SHEET_STATE = {
    "JABALPUR REGION": "Madhya Pradesh",
    "UP REGION": "Uttar Pradesh",
    "BPL, INDR & GWALIOR REGION": "Madhya Pradesh",
}

SHEET_CITY_FALLBACK = {
    "JABALPUR REGION": "Jabalpur Region",
    "UP REGION": "Uttar Pradesh Region",
    "BPL, INDR & GWALIOR REGION": "Bhopal, Indore & Gwalior Region",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_name(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "", clean(value).upper())


def normalize_address(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "", clean(value).upper())


def valid_email(value: Any) -> str | None:
    candidate = clean(value).lower()
    return candidate if candidate.upper() not in PLACEHOLDERS and EMAIL_RE.fullmatch(candidate) else None


def valid_gstin(value: Any) -> str | None:
    candidate = re.sub(r"\s+", "", clean(value).upper())
    return candidate if GSTIN_RE.fullmatch(candidate) else None


def normalize_phones(value: Any) -> list[str]:
    raw = clean(value)
    phones: list[str] = []
    for part in re.split(r"[,;/]", raw):
        digits = re.sub(r"\D", "", part)
        if len(digits) == 11 and digits.startswith("0"):
            digits = digits[1:]
        elif len(digits) == 12 and digits.startswith("91"):
            digits = digits[2:]
        elif len(digits) == 13 and digits.startswith("091"):
            digits = digits[3:]
        if len(digits) == 10 and digits[0] in "6789":
            phone = f"+91{digits}"
            if phone not in phones:
                phones.append(phone)
    return phones


def city_from_address(address: str, sheet: str, state: str) -> str:
    if not address:
        return SHEET_CITY_FALLBACK[sheet]
    parts = [part.strip(" .-") for part in re.split(r"[,\n]", address) if part.strip(" .-")]
    state_key = normalize_name(state)
    for part in reversed(parts):
        candidate = PINCODE_RE.sub("", part).strip(" .-")
        normalized = normalize_name(candidate)
        if not candidate or normalized == state_key or candidate.isdigit():
            continue
        if len(candidate) <= 60:
            return candidate
    return SHEET_CITY_FALLBACK[sheet]


def choose_text(values: list[str]) -> str:
    usable = [value for value in values if value and value.upper() not in PLACEHOLDERS]
    return max(usable, key=lambda value: (len(value), value)) if usable else ""


def extract(source: Path) -> dict[str, Any]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    raw_records: list[dict[str, Any]] = []

    def add_record(sheet: str, row: int, values: list[Any]) -> None:
        name, email, phone, address, gstin = values
        cleaned_name = clean(name)
        if not cleaned_name or cleaned_name.upper() in PLACEHOLDERS:
            return
        cleaned_address = clean(address)
        if cleaned_address.upper() in PLACEHOLDERS:
            cleaned_address = ""
        raw_records.append(
            {
                "sheet": sheet,
                "row": row,
                "name": cleaned_name,
                "email": valid_email(email),
                "rawPhone": clean(phone),
                "phones": normalize_phones(phone),
                "address": cleaned_address,
                "gstin": valid_gstin(gstin),
                "state": SHEET_STATE[sheet],
            }
        )

    for worksheet in workbook.worksheets:
        sheet = worksheet.title.strip()
        if sheet not in SHEET_STATE:
            continue
        if sheet == "UP REGION":
            # The first UP party is shifted into F:J beside the A:E headers.
            add_record(sheet, 1, [worksheet.cell(1, column).value for column in range(6, 11)])
            for row in range(2, worksheet.max_row + 1):
                add_record(sheet, row, [worksheet.cell(row, column).value for column in range(1, 6)])
        else:
            for row in range(2, worksheet.max_row + 1):
                values = [worksheet.cell(row, column).value for column in range(1, worksheet.max_column + 1)]
                if worksheet.max_column == 4:
                    values = [values[0], None, values[1], values[2], values[3]]
                add_record(sheet, row, values[:5])

    importable = [record for record in raw_records if record["phones"]]
    skipped = [record for record in raw_records if not record["phones"]]
    grouped: "OrderedDict[str, list[dict[str, Any]]]" = OrderedDict()
    for record in importable:
        if record["gstin"]:
            key = f"GSTIN:{record['gstin']}"
        else:
            phone_key = ",".join(sorted(record["phones"]))
            key = ":".join(
                ["PARTY", record["sheet"], normalize_name(record["name"]), phone_key, normalize_address(record["address"])]
            )
        grouped.setdefault(key, []).append(record)

    parties: list[dict[str, Any]] = []
    for key, records in grouped.items():
        phones: list[str] = []
        for record in records:
            for phone in record["phones"]:
                if phone not in phones:
                    phones.append(phone)
        name = choose_text([record["name"] for record in records])
        address = choose_text([record["address"] for record in records])
        state = records[0]["state"]
        pincode_match = PINCODE_RE.search(address)
        parties.append(
            {
                "sourceKey": key,
                "sourceRows": [f"{record['sheet']}!{record['row']}" for record in records],
                "region": records[0]["sheet"],
                "name": name,
                "phones": phones,
                "email": next((record["email"] for record in records if record["email"]), None),
                "address": address or "Address not provided",
                "city": city_from_address(address, records[0]["sheet"], state),
                "state": state,
                "pincode": pincode_match.group(1) if pincode_match else "",
                "gstin": next((record["gstin"] for record in records if record["gstin"]), None),
            }
        )

    unique_phones = {phone for party in parties for phone in party["phones"]}
    return {
        "sourceFile": source.name,
        "stats": {
            "rawPartyRows": len(raw_records),
            "importableRows": len(importable),
            "skippedRowsWithoutValidPhone": len(skipped),
            "distinctParties": len(parties),
            "uniquePhones": len(unique_phones),
            "multiPhoneRows": sum(len(record["phones"]) > 1 for record in raw_records),
            "mergedDuplicateRows": len(importable) - len(parties),
        },
        "parties": parties,
        "skippedRows": [
            {
                "sourceRow": f"{record['sheet']}!{record['row']}",
                "name": record["name"],
                "rawPhone": record["rawPhone"],
            }
            for record in skipped
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    result = extract(args.source)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["stats"], indent=2))


if __name__ == "__main__":
    main()
